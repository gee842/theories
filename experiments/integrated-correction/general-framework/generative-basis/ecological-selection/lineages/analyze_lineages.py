"""Source-bound genealogy and linked life-cycle assay reconstruction.

Read-only openpyxl extraction. No author notebook is executed. No Bayesian
ancestry/mutation inference or biological-replication inference is performed.
"""
from collections import Counter, defaultdict
import csv
from io import BytesIO
import itertools
import json
from pathlib import Path
from zipfile import ZipFile

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = next(p for p in HERE.parents if (p/'RESEARCH.md').exists())
SOURCE = ROOT/'sources/evolutionary-selection/doulcier-2025'
ARCHIVE = SOURCE/'lce-genealogies-code-and-data.zip'
PREFIX = 'lce-genealogies-main/'
COUNTS = PREFIX+'raw/19-09-27/Counts_final.xlsx'
ASSAYS = PREFIX+'raw/24-04-03/Lifecycle fitness tests.xlsx'


def read(book, sheet):
    values = list(book[sheet].values)
    return [dict(zip(values[0], row), source_row=i) for i,row in enumerate(values[1:],2)]


def key(t,tube):
    return f'{t}:{tube}'


def positive(x):
    return isinstance(x,(int,float)) and x > 0


def main():
    with ZipFile(ARCHIVE) as z:
        counts = openpyxl.load_workbook(BytesIO(z.read(COUNTS)),read_only=True,data_only=True)
        assays = openpyxl.load_workbook(BytesIO(z.read(ASSAYS)),read_only=True,data_only=True)
    treatments, exported, omitted = {}, [], []
    for sheet in counts.sheetnames:
        allrows=read(counts,sheet)
        assert len(allrows)==960
        index={}
        for r in allrows:
            if r['Condition']=='MISS':
                omitted.append(dict(sheet=sheet,**r));continue
            t=2*r['Cycle']+r['Phase'];tube=r['Rack']+str(r['Lineage'])
            assert (t,tube) not in index
            index[t,tube]=r
        ancestry={}; parents={}; missing=[]; nonviable=[]
        phase_counts=defaultdict(Counter)
        for (t,tube),r in sorted(index.items()):
            parent=None
            if t>1:
                prior=index.get((t-1,tube))
                if prior is not None:
                    donor=prior['Replaced'] if prior['Replaced'] not in (0,None,'na') else tube
                    if (t-1,donor) in index:
                        parent=key(t-1,donor)
                        if index[t-1,donor]['Extinct']!=0:nonviable.append([key(t,tube),parent])
                if parent is None:missing.append(key(t,tube))
            name=key(t,tube);parents[name]=parent
            ancestry[name]={name}|(ancestry[parent] if parent else set())
            assert r['Extinct'] in (0,1)
            phase_counts[t]['observed']+=1
            phase_counts[t]['survived']+=int(r['Extinct']==0)
            phase_counts[t]['flagged']+=int(r['Condition']!=0)
            exported.append(dict(sheet=sheet,source_row=r['source_row'],node=name,parent=parent or '',
                                 cycle=r['Cycle']+1,phase=r['Phase'],tube=tube,mat=r['Mat'],
                                 count_ws=r['Count WS'],count_sm=r['Count SM'],extinct=r['Extinct'],
                                 replaced=r['Replaced'],condition=r['Condition']))
        terminal=[key(t,tube) for (t,tube),r in index.items() if t==20 and r['Extinct']==0]
        founder_counts=Counter()
        for node in terminal:
            roots=[n for n in ancestry[node] if n.startswith('1:')]
            assert len(roots)<=1
            founder_counts[roots[0] if roots else 'unresolved']+=1
        common=set.intersection(*(ancestry[n] for n in terminal))
        treatments[sheet]=dict(raw_rows=960,retained_nodes=len(index),miss_rows=960-len(index),
                              source_condition_counts=dict(Counter(str(r['Condition']) for r in allrows)),
                              unresolved_parent_nodes=missing,nonviable_donors=nonviable,
                              observed_parent_edges=sum(v is not None for v in parents.values()),
                              terminal_viable_nodes=len(terminal),terminal_founder_counts=dict(sorted(founder_counts.items())),
                              terminal_common_ancestors=sorted(common,key=lambda n:(int(n.split(':')[0]),n)),
                              descendants_of_2_II_A5=sum('4:A5' in ancestry[n] for n in terminal),
                              phase_counts={str(t):dict(v) for t,v in sorted(phase_counts.items())})
    with (HERE/'genealogy-reading.csv').open('w',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(exported[0]));writer.writeheader();writer.writerows(exported)
    (HERE/'missing-reading.json').write_text(json.dumps(omitted,indent=2)+'\n')
    rows=read(assays,'raw')
    focal={}; assay_export=[]
    for strain in ['PE-2-WT-A5','PE-2-WT-A5 mutL']:
        rs=[r for r in rows if r['Strain']==strain and str(r['Size']).lower()=='large'
            and str(r['Date']).startswith('2020-06-02')]
        assert all(r['Condition']==0 and r['Replaced by']=='NA' for r in rs)
        phases={p:{(r['Rack'],r['Tube']):r for r in rs if r['Phase']==p} for p in (1,2)}
        assert len(phases[1])==72
        assert set(phases[2])=={k for k,r in phases[1].items() if r['Extinct']==0}
        assert all((positive(r['Mat']) and positive(r['Count SM']))==(r['Extinct']==0) for r in phases[1].values())
        assert all(positive(r['Count WS'])==(r['Extinct']==0) for r in phases[2].values())
        by_rack=[]
        for rack in sorted({k[0] for k in phases[1]}):
            start=[k for k in phases[1] if k[0]==rack]
            through=[k for k in start if k in phases[2]]
            final=[k for k in through if phases[2][k]['Extinct']==0]
            by_rack.append(dict(rack=rack,start=len(start),phase1_survived=len(through),cycle_completed=len(final)))
        n1=len(phases[1]);n2=len(phases[2]);end=sum(r['Extinct']==0 for r in phases[2].values())
        focal[strain]=dict(initial=n1,phase1_survived=n2,phase2_survived=end,
                           cycle_fraction=end/n1,phase2_conditional_fraction=end/n2,
                           racks=by_rack,joined_by='same date, strain, size, rack and tube; no replacement recorded')
        for r in rs:assay_export.append({**r,'Date':str(r['Date'])})
    with (HERE/'linked-assay-reading.csv').open('w',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(assay_export[0]));writer.writeheader();writer.writerows(assay_export)
    a,b=[focal[x] for x in ['PE-2-WT-A5','PE-2-WT-A5 mutL']]
    differences=[(a['phase2_survived']-x['cycle_completed'])/64-(b['phase2_survived']-y['cycle_completed'])/64
                 for x,y in itertools.product(a['racks'],b['racks'])]
    # Separate audit of the author's product of phase-I marginal frequencies.
    ancestral=[r for r in rows if r['Strain']=='Line-17 mutS' and str(r['Size']).lower()=='large' and r['Phase']==1]
    known=Counter();unknown=[]
    for r in ancestral:
        if not isinstance(r['Count SM'],(int,float)):
            unknown.append({k:str(r[k]) for k in ['source_row','Date','Mat','Count SM','Extinct']});continue
        known[f'{int(positive(r["Mat"]))}{int(positive(r["Count SM"]))}']+=1
    assert len(ancestral)==184 and len(unknown)==1 and unknown[0]['Mat']=='1'
    joint=known['11'];mat=known['11']+known['10'];sm=known['11']+known['01'];n=len(ancestral)
    complete=sum(known.values())
    result=dict(treatments=treatments,total_raw_genealogy_rows=3840,known_genealogy_nodes=len(exported),
                source_miss_rows=len(omitted),focal_assay=focal,
                matched_date_risk_difference=a['cycle_fraction']-b['cycle_fraction'],
                matched_date_ratio=a['cycle_fraction']/b['cycle_fraction'],
                delete_one_rack_each_difference_range=[min(differences),max(differences)],
                rack_sensitivity_scope='81 deletion combinations; descriptive, not confidence intervals or independently evolved mutations.',
                ancestral_phase1=dict(n=n,complete_n=complete,known_table=dict(known),unknown_sm=unknown,
                                     complete_mat_success=mat,sm_known_success=sm,joint_known_success=joint,
                                     complete_joint_fraction=joint/complete,
                                     complete_marginal_product=mat*sm/complete**2,
                                     full_record_joint_range=[joint/n,(joint+1)/n],
                                     full_record_marginal_product_range=[(mat+1)*sm/n**2,(mat+1)*(sm+1)/n**2],
                                     observed_extant=sum(r['Extinct']==0 for r in ancestral)),
                limits=['No imputed parents across missing phases.',
                        'Condition flags are retained; raw recorded parent links are not genotype proofs.',
                        'No Bayesian mutation propagation or sequencing alignment rerun.',
                        'Single mutation origin and one matched-date comparison; no independent evolution replication.',
                        'Growth medium, transfer schedule and replacement lottery remain experimentally supplied.'])
    (HERE/'lineage-results.json').write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps({k:result[k] for k in ['known_genealogy_nodes','source_miss_rows','focal_assay','matched_date_risk_difference','matched_date_ratio','delete_one_rack_each_difference_range','ancestral_phase1']},indent=2))


if __name__=='__main__':main()
