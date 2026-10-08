#!/usr/bin/env python3
"""Source-bound summaries and a scoped replay of Karita's four-state model.

Reads workbooks with openpyxl and independently checks numeric areas and the
literal pellicle count cells through raw XLSX XML. No biological replication.
Only inspected notebook parameter/function cells are executed; no outputs or
plotting cells are trusted as results. All numerical replay is illustrative.
"""
from collections import defaultdict
from datetime import datetime, timedelta
from fractions import Fraction
from pathlib import Path
import csv
import hashlib
import json
import math
import statistics
import xml.etree.ElementTree as ET
import zipfile
import openpyxl
from scipy.integrate import solve_ivp

HERE = Path(__file__).resolve().parent
ROOT = next(p for p in HERE.parents if (p / 'RESEARCH.md').exists())
SOURCE = ROOT / 'sources/evolutionary-selection/karita-2026'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def raw_cells(path):
    """Independent XML path for the first (and only) sheet of two workbooks."""
    with zipfile.ZipFile(path) as z:
        strings = []
        if 'xl/sharedStrings.xml' in z.namelist():
            for item in ET.fromstring(z.read('xl/sharedStrings.xml')).findall('s:si', NS):
                strings.append(''.join(t.text or '' for t in item.findall('.//s:t', NS)))
        cells = {}
        for cell in ET.fromstring(z.read('xl/worksheets/sheet1.xml')).findall('.//s:c', NS):
            value = cell.find('s:v', NS)
            if cell.get('t') == 'inlineStr':
                cells[cell.get('r')] = ''.join(t.text or '' for t in cell.findall('.//s:t', NS))
            elif value is not None and value.text is not None:
                cells[cell.get('r')] = strings[int(value.text)] if cell.get('t') == 's' else Fraction(value.text)
        return cells


def derivative(x, p):
    """Independent flux decomposition, ordered sA,sB,wA,wB."""
    sa, sb, wa, wb = x
    gsa, gsb, gwa, gwb = p['r_sa']*sa, p['r_sb']*sb, p['r_wa']*wa, p['r_wb']*wb
    if sa + wa >= p['K_a']:
        weight_s = p['d_s'] * sa * p['switch']
        weight_w = p['d_w'] * wa
        ds = weight_s/(weight_s+weight_w)*(gsa+gwa)
        dw = weight_w/(weight_s+weight_w)*(gsa+gwa)
    else:
        ds, dw = p['d_s']*sa, p['d_w']*wa
    vacancy = 1 - (sa+wa)/p['K_a']
    ms, mw = p['m_s']*sb*vacancy, p['m_w']*wb*vacancy
    qa, qb = p['mu']*(gsa-gwa), p['mu']*(gsb-gwb)
    return [gsa-ds+ms-qa, gsb+ds-ms-qb, gwa-dw+mw+qa, gwb+dw-mw+qb]


def integrate(p, step, horizon=96):
    x = [100000., 10000000., 0., 0.]
    record = []
    overshoot = 0.
    minimum = 0.
    for n in range(round(horizon/step)+1):
        if n % round(12/step) == 0:
            record.append({'source_time_units': n*step, 'source_nominal_days': n*step/12,
                           'sA': x[0], 'sB': x[1], 'wA': x[2], 'wB': x[3],
                           'surface_WS_fraction': x[2]/(x[0]+x[2])})
        overshoot = max(overshoot, (x[0]+x[2])/p['K_a']-1)
        minimum = min(minimum, *x)
        if n < round(horizon/step):
            dx = derivative(x, p)
            x = [v+step*d for v, d in zip(x, dx)]
    return {'step': step, 'records': record, 'max_relative_capacity_overshoot': overshoot,
            'minimum_population': minimum}


def event_integrate(p, tolerance, horizon=96):
    """Locate capacity, then evolve the exact saturated coordinate q=wA/K.

    This is an explicitly identified numerical adaptation. The notebook's
    overshot branch is not silently clipped and no negative rate is zeroed.
    Normalized abundances make the tolerances comparable across components.
    """
    capacity = p['K_a']
    def initial_rhs(t, y):
        sa, sb, wa, wb = y
        vacancy = 1-sa-wa
        return [p['r_sa']*sa-p['d_s']*sa+p['m_s']*sb*vacancy-p['mu']*(p['r_sa']*sa-p['r_wa']*wa),
                p['r_sb']*sb+p['d_s']*sa-p['m_s']*sb*vacancy-p['mu']*(p['r_sb']*sb-p['r_wb']*wb),
                p['r_wa']*wa-p['d_w']*wa+p['m_w']*wb*vacancy+p['mu']*(p['r_sa']*sa-p['r_wa']*wa),
                p['r_wb']*wb+p['d_w']*wa-p['m_w']*wb*vacancy+p['mu']*(p['r_sb']*sb-p['r_wb']*wb)]
    def cap_event(t,y): return y[0]+y[2]-1
    cap_event.terminal = True
    cap_event.direction = 1
    initial = solve_ivp(initial_rhs, [0,horizon], [1e5/capacity,1e7/capacity,0,0],
                        events=cap_event, rtol=tolerance, atol=tolerance*1e-8,
                        dense_output=True, max_step=.1)
    assert initial.success and len(initial.t_events[0]) == 1
    hit = float(initial.t_events[0][0])
    sa,sb,wa,wb = initial.y_events[0][0]
    assert abs(sa+wa-1)<1e-12
    def saturated_rhs(t,y):
        q, sb, wb = y
        sa,wa = 1-q,q
        gs,gw = p['r_sa']*sa,p['r_wa']*wa
        denom = p['d_s']*sa*p['switch']+p['d_w']*wa
        ds = p['d_s']*sa*p['switch']/denom*(gs+gw)
        dw = p['d_w']*wa/denom*(gs+gw)
        ma,mb = p['mu']*(gs-gw),p['mu']*(p['r_sb']*sb-p['r_wb']*wb)
        return [gw-dw+ma,p['r_sb']*sb+ds-mb,p['r_wb']*wb+dw+mb]
    later = solve_ivp(saturated_rhs,[hit,horizon],[wa,sb,wb],rtol=tolerance,
                      atol=tolerance*1e-8,dense_output=True,max_step=.1)
    assert later.success
    records=[]
    for t in range(0,horizon+1,12):
        if t<hit:
            sa,sb,wa,wb=initial.sol(t)
        else:
            q,sb,wb=later.sol(t);sa,wa=1-q,q
        assert min(sa,sb,wa,wb)>=0 and 0<=wa/(sa+wa)<=1
        records.append({'source_time_units':t,'source_nominal_days':t/12,
                        'sA':float(sa*capacity),'sB':float(sb*capacity),
                        'wA':float(wa*capacity),'wB':float(wb*capacity),
                        'surface_WS_fraction':float(wa/(sa+wa))})
    return {'relative_tolerance':tolerance,'capacity_event_time':hit,'records':records,
            'nonnegative_at_solver_nodes':bool((initial.y>=0).all() and (later.y>=0).all() and (later.y[0]<=1).all()),
            'scope':'Numerical event-based adaptation with exact saturated coordinate, not a certified ODE solution or biological prediction.'}


def main():
    manifest = json.loads((SOURCE / 'manifest.json').read_text())
    for f in manifest['files']:
        p = SOURCE / f['file']
        assert p.stat().st_size == f['bytes'] and sha(p) == f['sha256']
        if 'publisher_md5' in f:
            assert hashlib.md5(p.read_bytes()).hexdigest() == f['publisher_md5']
    sheet = openpyxl.load_workbook(SOURCE / 'Pellicle_success.xlsx', data_only=True).active
    raw = raw_cells(SOURCE / 'Pellicle_success.xlsx')
    cohorts = []
    grouped = defaultdict(lambda: [0, 0])
    for col in [1, 4, 7, 10]:
        date = sheet.cell(2, col+1).value.isoformat()[:10]
        serial = raw[sheet.cell(2, col+1).coordinate]
        assert (datetime(1899,12,30)+timedelta(days=float(serial))).isoformat()[:10] == date
        for row in range(4, sheet.max_row+1):
            d, count = sheet.cell(row, col).value, sheet.cell(row, col+1).value
            if d is None:
                assert count is None
                continue
            assert raw[sheet.cell(row, col).coordinate] == Fraction(str(d))
            assert raw[sheet.cell(row, col+1).coordinate] == count
            success, total = map(int, count.split(','))
            assert 0 <= success <= total
            cohorts.append({'date': date, 'sheet': sheet.title, 'count_cell': sheet.cell(row, col+1).coordinate,
                            'log10_dilution_factor': -d, 'success': success, 'total': total})
            grouped[-d][0] += success
            grouped[-d][1] += total
    with (HERE / 'pellicle-cohorts.csv').open('w') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(cohorts[0])); writer.writeheader(); writer.writerows(cohorts)
    assert dict(grouped) == {2.: [3,3], 2.5: [6,6], 3.: [12,12], 3.5: [10,15], 4.: [0,9]}
    size_sheet = openpyxl.load_workbook(SOURCE / 'Collapsing_colony_size.xlsx', data_only=True).active
    areas = [(i, values[0]) for i, values in enumerate(size_sheet.values, 1) if isinstance(values[0], (int, float))]
    xml_areas = [(int(c[1:]), v) for c,v in raw_cells(SOURCE / 'Collapsing_colony_size.xlsx').items()
                 if c.startswith('A') and c[1:].isdigit() and isinstance(v, Fraction)]
    # Excel numeric cells store decimal spellings of binary floating values.
    # Compare the decoded doubles, preserving the unchanged XML in the source.
    assert areas == [(i, float(v)) for i,v in xml_areas]
    diameters = [2*math.sqrt(v/math.pi) for _,v in areas]
    with (HERE / 'colony-areas.csv').open('w') as handle:
        writer = csv.writer(handle); writer.writerow(['source_row','source_area','circle_equivalent_diameter'])
        writer.writerows((i, a, d) for (i,a),d in zip(areas,diameters))
    notebook = json.loads((SOURCE / '4state-model.ipynb').read_text())
    namespace = {'math': math}
    exec(''.join(notebook['cells'][1]['source']), namespace)
    exec(''.join(notebook['cells'][2]['source']), namespace)
    keys = ['r_sa','r_wa','r_sb','r_wb','d_s','d_w','switch','m_s','m_w','mu','K_a']
    params = {k: namespace[k] for k in keys}
    states = [[1e5,1e7,0.,0.],[1e8,1e9,2e7,3e6],[6e9,1e10,4e9,2e10],
              [6.006e9,1e10,4e9,2e10],[0.,1e10,1e10,1e10]]
    error = 0.
    for x in states:
        expected = [namespace[k](*x) for k in ['d_sa','d_sb','d_wa','d_wb']]
        actual = derivative(x, params)
        error = max(error, *(abs(a-b)/max(1,abs(a),abs(b)) for a,b in zip(expected,actual)))
        assert all(math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-5) for a,b in zip(expected,actual))
    scenarios = {}
    event_scenarios = {}
    for name, dw in [('source', .001), ('equal_dispersal', .002), ('neutral_retention_slice', .0017)]:
        p = dict(params, d_w=dw)
        scenarios[name] = [integrate(p, h) for h in [.001,.0005]]
        event_scenarios[name] = [event_integrate(p, tol) for tol in [1e-8,1e-10]]
        assert all(r['nonnegative_at_solver_nodes'] for r in event_scenarios[name])
    # Exact rational check of the specified local balance, independent of floats.
    F = Fraction
    s=w=F(1);rs=F(1);rw=F(17,20);ds=F(1,500);dw=F(1,1000)
    net = rw*w-dw*w/(ds*s+dw*w)*(rs*s+rw*w)
    assert net == F(7,30)
    result = {
        'scope': 'Empirical extraction and illustrative conventional-model replay; no biological replication, fitted kinetic prediction, ODE proof or comparative superiority.',
        'checks_passed': True,
        'pellicle_cohorts': len(cohorts), 'pellicle_cultures': sum(c['total'] for c in cohorts),
        'pellicle_pooled': [{'log10_dilution_factor': d, 'success': v[0], 'total': v[1]} for d,v in sorted(grouped.items())],
        'same_dilution_3_5_cohorts': [c for c in cohorts if c['log10_dilution_factor']==3.5],
        'colony_areas': {'raw_n': len(areas), 'paper_caption_n': 769,
                         'mean_diameter': statistics.mean(diameters), 'median_diameter': statistics.median(diameters),
                         'range': [min(diameters),max(diameters)],
                         'n_le_65': sum(d<=65 for d in diameters),
                         'cutoff_is_not_a_verified_source_exclusion_rule': True,
                         'culture_nesting_unresolved': True},
        'notebook_parameters': params, 'source_function_comparisons': len(states)*4,
        'max_relative_derivative_difference': error,
        'source_nominal_time_conversion_not_empirical_calibration': True,
        'model_scenarios': scenarios,
        'event_model_scenarios': event_scenarios,
        'source_same_slice_rational_net': str(net),
        'raw_xml_crosscheck': {'literal_cohort_cells': len(cohorts), 'areas': len(areas)},
        'source_manifest_sha256': sha(SOURCE / 'manifest.json'),
        'script_sha256': sha(Path(__file__)),
        'outputs_sha256': {n:sha(HERE/n) for n in ['pellicle-cohorts.csv','colony-areas.csv']},
    }
    (HERE / 'results.json').write_text(json.dumps(result, indent=2)+'\n')
    print(json.dumps({k:result[k] for k in ['checks_passed','pellicle_cohorts','pellicle_cultures','colony_areas','source_function_comparisons','max_relative_derivative_difference']},indent=2))
    for name, runs in scenarios.items():
        print(name, [(r['step'],r['records'][-1]['surface_WS_fraction'],r['max_relative_capacity_overshoot']) for r in runs])
        print('event',[(r['relative_tolerance'],r['records'][-1]['surface_WS_fraction'],r['capacity_event_time']) for r in event_scenarios[name]])


if __name__ == '__main__':
    main()
