"""Read-only reconstruction of Driscoll & Travisano 2017 Data 5.

No individual-particle measurements or biological replicates are inferred from
video means. All contrasts are descriptive. Run with openpyxl available.
"""
from collections import defaultdict
import csv
import hashlib
import itertools
import json
from pathlib import Path
import statistics

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = next(p for p in HERE.parents if (p / "RESEARCH.md").exists())
SOURCE = ROOT / "sources/evolutionary-selection/driscoll-2017"
BOOK = SOURCE / "41467_2017_BFncomms15707_MOESM80_ESM.xlsx"


def main():
    book = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    sheet = book["Mean cluster velocity"]
    expected = ["Species", "Media", "Density (%)", "Replicate", "Video",
                "Mean velocity (mm / min)"]
    assert [sheet.cell(1, i).value for i in range(1, 7)] == expected
    records = []
    for row in range(2, 56):
        sp, medium, density, rep, video, speed = [sheet.cell(row, i).value for i in range(1, 7)]
        assert sp in ("K. lactis", "S. cerevisiae")
        assert rep in (1, 2, 3) and video in (1, 2, 3)
        records.append(dict(source_row=row, species=sp, medium=medium,
                            density_percent=density, technical_replicate=rep,
                            video=video, mean_velocity_mm_min=speed))
    assert len(records) == 54
    assert len({(r['species'], r['medium'], r['density_percent'],
                 r['technical_replicate'], r['video']) for r in records}) == 54
    with (HERE / "video-means.csv").open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    grouped = defaultdict(lambda: defaultdict(list))
    for r in records:
        key = f"{r['species']}|{r['medium']}|{r['density_percent']}"
        grouped[key][r['technical_replicate']].append(r['mean_velocity_mm_min'])
    groups = {}
    for key, reps in sorted(grouped.items()):
        assert sorted(reps) == [1, 2, 3]
        assert all(len(v) == 3 for v in reps.values())
        means = [statistics.mean(reps[i]) for i in (1, 2, 3)]
        groups[key] = dict(video_count=9, technical_replicate_count=3,
                           replicate_means_mm_min=means,
                           mean_mm_min=statistics.mean(means),
                           replicate_mean_range_mm_min=[min(means), max(means)],
                           replicate_mean_sd_mm_min=statistics.stdev(means),
                           leave_one_replicate_out_means_mm_min=[
                               statistics.mean(means[:i] + means[i+1:]) for i in range(3)])
    def mean(key):
        return groups[key]['mean_mm_min']
    k_low, k_high = "K. lactis|YPD|5", "K. lactis|YPD|100"
    s_low, s_high = "S. cerevisiae|YPD|5", "S. cerevisiae|YPD|100"
    gal, dex = "K. lactis|Galactose|100", "K. lactis|Dextrose|100"
    contrasts = {}
    for name, high, low in [('K_lactis_density', k_high, k_low),
                            ('S_cerevisiae_density', s_high, s_low),
                            ('K_lactis_sugar', dex, gal)]:
        a, b = groups[high], groups[low]
        # Cross-treatment replicate labels do NOT imply matched cultures.
        delete_one = [x-y for x, y in itertools.product(
            a['leave_one_replicate_out_means_mm_min'],
            b['leave_one_replicate_out_means_mm_min'])]
        contrasts[name] = dict(high=high, low=low,
                              mean_difference_mm_min=mean(high)-mean(low),
                              ratio_of_means=mean(high)/mean(low),
                              min_high_rep_minus_max_low_rep_mm_min=(
                                  min(a['replicate_means_mm_min'])-max(b['replicate_means_mm_min'])),
                              delete_one_each_condition_difference_range_mm_min=[min(delete_one), max(delete_one)])
    interaction = (mean(k_high)-mean(k_low))-(mean(s_high)-mean(s_low))
    deletion_interactions = [a-b-c+d for a,b,c,d in itertools.product(*[
        groups[key]['leave_one_replicate_out_means_mm_min']
        for key in (k_high, k_low, s_high, s_low)])]
    result = dict(
        source_sha256=hashlib.sha256(BOOK.read_bytes()).hexdigest(),
        source_sheet=sheet.title, original_rows=[2, 55],
        observed_unit='mean cluster speed within a video',
        nesting='54 videos in 18 technical replicates, six conditions; 3 videos/replicate',
        biological_replication='Not established by these video data; methods call the replicate IDs technical.',
        units='mm/min', groups=groups, contrasts=contrasts,
        species_by_density_difference_mm_min=interaction,
        delete_one_each_condition_interaction_range_mm_min=[min(deletion_interactions), max(deletion_interactions)],
        sensitivity_scope='All 81 combinations of deleting one technical replicate in each YPD condition. Not a confidence interval or a population inference.',
        limits=['No particle-level size, neighbor, or trajectory records in this worksheet.',
                'No source-matched material/geometry/velocity/offspring history.',
                'No new hydrodynamic fit, independent biological replication, or superiority claim.'])
    (HERE / 'interaction-results.json').write_text(json.dumps(result, indent=2)+'\n')
    print(json.dumps({'groups':groups,'contrasts':contrasts,
                      'species_by_density_difference_mm_min':interaction,
                      'deletion_interaction_range':result['delete_one_each_condition_interaction_range_mm_min']},indent=2))


if __name__ == '__main__':
    main()
