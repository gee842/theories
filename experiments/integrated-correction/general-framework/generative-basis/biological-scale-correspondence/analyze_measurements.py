#!/usr/bin/env python3
"""Descriptive reanalysis of published data, not a fit of the Lean model.
No biological observations are joined across figures. Blank cells stay missing.
Cross-check the JS extraction against read-only openpyxl and means against numpy.
"""
from pathlib import Path
import json, hashlib, statistics
from datetime import datetime, timezone
import openpyxl
import numpy as np
ROOT=Path(__file__).resolve().parent
REPO=next(p for p in ROOT.parents if (p/'RESEARCH.md').exists())
SOURCE=REPO/'sources/biological-scale-correspondence/elife-65935'

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def col(n):
    s=''
    while n:
        n,r=divmod(n-1,26);s=chr(65+r)+s
    return s

def load(fig):
    p=SOURCE/f'elife-65935-fig{fig}-data1-v2.xlsx'
    extract=ROOT/'extracted'/f'figure{fig}.json'
    d=json.loads(extract.read_text())
    wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
    assert wb.sheetnames==[d['sheet']]
    ws=wb[d['sheet']]
    actual=[list(r) for r in ws.iter_rows(values_only=True)]
    expected=d['values']
    assert len(actual)==len(expected)
    assert all(a==b for a,b in zip(actual,expected))
    wb.close()
    return d

def summarize(d,row,start,end,**labels):
    values=[];cells=[]
    for c in range(start,end):
        value=d['values'][row][c]
        if value is None: continue
        assert isinstance(value,(int,float)) and not isinstance(value,bool)
        assert np.isfinite(value)
        values.append(value);cells.append(f'{col(c+1)}{row+1}')
    mean=statistics.mean(values);sd=statistics.stdev(values)
    assert abs(mean-float(np.mean(values)))<1e-12
    assert abs(sd-float(np.std(values,ddof=1)))<1e-12
    return dict(**labels,n=len(values),mean=mean,sample_sd=sd,cells=cells)

f2=load(2);f4=load(4)
assert f2['values'][0][0]=='Anterior' and f2['values'][7][0]=='Medial'
assert [f2['values'][1][i] for i in (1,19,37)]==['E. coli','P. aerugonisa','E. faecalis']
assert f4['values'][0][1]=='WT' and f4['values'][0][10]=='npr-1(ad609)' and f4['values'][0][19]=='npr-1(ok1447)'
regions=[];occupancy=[]
for region,rows in [('Anterior',[2,3,4]),('Medial',[9,10,11])]:
 for r in rows:
  for species,start,end in [('E. coli',1,19),('P. aeruginosa',19,37),('E. faecalis',37,55)]:
   regions.append(summarize(f2,r,start,end,region=region,hours=f2['values'][r][0],species=species,unit='micrometres'))
for r in range(1,6):
 for strain,start,end in [('WT',1,10),('npr-1(ad609)',10,19),('npr-1(ok1447)',19,28)]:
  row=summarize(f4,r,start,end,strain=strain,hours=f4['values'][r][0],unit='occupancy fraction',oxygen='8 percent' if r<5 else 'atmospheric after chamber removal')
  assert row['n']==9 and all(0<=v<=1 for v in f4['values'][r][start:end])
  occupancy.append(row)
changes=[]
for strain in ['WT','npr-1(ad609)','npr-1(ok1447)']:
 a=next(x for x in occupancy if x['strain']==strain and x['hours']==4)
 b=next(x for x in occupancy if x['strain']==strain and x['hours']==5)
 changes.append(dict(strain=strain,mean_4h=a['mean'],mean_5h=b['mean'],difference_of_means=b['mean']-a['mean']))
manifest={'retrieved_utc':datetime.now(timezone.utc).isoformat(),'article_doi':'10.7554/eLife.65935','authors':'Adam Filipowicz, Jonathan Lalsiamthara, Alejandro Aballay','year':2021,'license':'CC BY 4.0 as stated by eLife; attribution retained','discovery':'Links obtained from https://api.elifesciences.org/articles/65935 and cross-identified with PMC source-data captions','files':[{'file':p.name,'url':f'https://cdn.elifesciences.org/articles/65935/{p.name}','bytes':p.stat().st_size,'sha256':sha(p)} for p in sorted(SOURCE.glob('*.xlsx'))]}
mp=SOURCE/'manifest.json'
if not mp.exists(): mp.write_text(json.dumps(manifest,indent=2)+'\n')
else:
 old=json.loads(mp.read_text());assert old['files']==manifest['files']
result={'status':'passed','scope':'Descriptive reanalysis, independent extraction agreement and numerical cross-check; not experimental replication or a calibrated mechanistic fit.','diameter_measurements':sum(x['n'] for x in regions),'occupancy_measurements':sum(x['n'] for x in occupancy),'diameter_groups':regions,'occupancy_groups':occupancy,'oxygen_transition':changes,'missingness':'Blank source cells excluded and not imputed. Different counts across regions preclude assumed individual pairing. No individual or plate identifiers join figures. Occupancy denominators are not supplied in this workbook.','source_files':manifest['files'],'extraction_checked_with':'openpyxl read_only/data_only against artifact-tool values; all cells equal','statistics_checked_with':'statistics against numpy within 1e-12; sample SD uses n-1','script_sha256':sha(Path(__file__)),'extractor_sha256':sha(ROOT/'extract_measurements.mjs')}
(ROOT/'results.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps({k:result[k] for k in ['status','diameter_measurements','occupancy_measurements','oxygen_transition']},indent=2))
print('anterior at 1 hour',[(x['species'],x['n'],x['mean']) for x in regions if x['region']=='Anterior' and x['hours']==1])
