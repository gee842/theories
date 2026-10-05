"""Independent, read-only raw-workbook audit; does not import primary code.

Run with the bundled Python from any directory. Writes only its own JSON report.
Means/variances use Fraction; Student-t uses density quadrature, not SciPy.
"""
from collections import Counter, defaultdict
from fractions import Fraction as Q
from hashlib import sha256
from math import exp, lgamma, pi, sqrt
from pathlib import Path
import json
import re

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
DATA = BASE / "data"
checks = []


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    checks.append(label)


def mean(seq):
    return sum(seq, Q(0)) / len(seq)


def variance(seq):
    mu = mean(seq)
    return sum(((z - mu) ** 2 for z in seq), Q(0)) / (len(seq) - 1)


def tcdf(x, df, panels=4096):
    """Simpson integration of Student-t density on [0,x], x >= 0."""
    normalizer = exp(lgamma((df + 1) / 2) - lgamma(df / 2)) / sqrt(df * pi)
    step = x / panels
    density = lambda v: normalizer * (1 + v * v / df) ** (-(df + 1) / 2)
    total = density(0) + density(x)
    for i in range(1, panels):
        total += (4 if i % 2 else 2) * density(i * step)
    return 0.5 + total * step / 3


source_names = ["Dataset2_Main.xlsx", "Dataset1_ChatCodes.xlsm", "Protocol_author.doc"]
before = {name: sha256((DATA / name).read_bytes()).hexdigest() for name in source_names}
wb = load_workbook(DATA / source_names[0], read_only=True, data_only=True)
raw = list(wb["Data"].iter_rows(min_row=2, max_row=289, values_only=True))
wb.close()
group_rows = defaultdict(list)
for row in raw:
    group_rows[int(row[7])].append(row)
check("288 participant rows and 72 groups", len(raw) == 288 and len(group_rows) == 72)
groups = {}
for g, rows in group_rows.items():
    check(f"group {g}: four rows", len(rows) == 4)
    # B, E, F, G, CD, CE, CF; all nonempty replicated values must agree.
    values = []
    for col in [1, 4, 5, 6, 81, 82, 83]:
        unique = {r[col] for r in rows if r[col] is not None}
        check(f"group {g}: column {col + 1} agrees", len(unique) == 1)
        values.append(unique.pop())
    sid, original, treatment, facil, e1, e2, e3 = values
    es = tuple(map(int, (e1, e2, e3)))
    check(f"group {g}: valid ES categories", all(z in range(4) for z in es))
    expected_original = {1: 3, 2: 2, 3: 4, 4: 1, 5: 5, 6: 6}[int(treatment)]
    check(f"group {g}: treatment mapping", original == expected_original and facil == (int(treatment) % 2 == 0))
    groups[g] = dict(session=str(sid).split("_", 1)[1], treatment=int(treatment), es=es,
                     y=int(es[1] >= 2 and es[2] >= 2))

session_groups = defaultdict(list)
for g, row in groups.items():
    session_groups[row["session"]].append(g)
check("31 true sessions", len(session_groups) == 31)
check("21 two-group and ten three-group sessions", Counter(map(len, session_groups.values())) == {2: 21, 3: 10})
check("one treatment per session", all(len({groups[g]["treatment"] for g in ids}) == 1 for ids in session_groups.values()))

chat = load_workbook(DATA / source_names[1], read_only=True, data_only=True)
headers = {}
for sheet in chat:
    hit = re.fullmatch(r"GROUP (\d+) \(([^)]+)\)", str(sheet.cell(1, 1).value))
    if hit:
        headers[int(hit[1])] = hit[2].replace(" ", "-")
check("all 72 chat headers confirm session suffixes", headers == {g: r["session"] for g, r in groups.items()})
coders = {}
for row in chat["Kappa (Interrater Reliability)"].iter_rows(min_row=2, max_row=73, values_only=True):
    coders[int(row[0])] = (tuple(map(int, row[7:10])), tuple(map(int, row[10:13])))
chat.close()
check("72 paired coder records", set(coders) == set(groups))
check("main codes exactly match coder 1", all(groups[g]["es"] == pair[0] for g, pair in coders.items()))
disagreements = [(g, i + 1, a[i], b[i]) for g, (a, b) in coders.items() for i in range(3) if a[i] != b[i]]
coder2y = {g: int(pair[1][1] >= 2 and pair[1][2] >= 2) for g, pair in coders.items()}
endpoint_disagreements = [g for g in groups if groups[g]["y"] != coder2y[g]]


def aggregate(outcomes):
    cells = {}
    for treatment in range(1, 7):
        ids = [g for g in groups if groups[g]["treatment"] == treatment]
        sessions = sorted({groups[g]["session"] for g in ids})
        sy = [mean([Q(outcomes[g]) for g in ids if groups[g]["session"] == sid]) for sid in sessions]
        cells[treatment] = dict(groups=len(ids), sessions=len(sessions), values=sy,
                                mean=mean(sy), variance=variance(sy), group_mean=mean([Q(outcomes[g]) for g in ids]))
    contrasts = [cells[j + 1]["mean"] - cells[j]["mean"] for j in [1, 3, 5]]
    group_contrasts = [cells[j + 1]["group_mean"] - cells[j]["group_mean"] for j in [1, 3, 5]]
    return cells, contrasts, group_contrasts


cells, contrasts, group_contrasts = aggregate({g: row["y"] for g, row in groups.items()})
components = [c["variance"] / c["sessions"] / 9 for c in cells.values()]
v = sum(components, Q(0))
df = float(v * v / sum((c * c / (cell["sessions"] - 1) for c, cell in zip(components, cells.values())), Q(0)))
effect, se = mean(contrasts), sqrt(float(v))
lo, hi = 0.0, 10.0
for _ in range(48):
    mid = (lo + hi) / 2
    if tcdf(mid, df) < 0.975:
        lo = mid
    else:
        hi = mid
critical = (lo + hi) / 2
ci = [float(effect) - critical * se, float(effect) + critical * se]
pvalue = 2 * (1 - tcdf(float(effect) / se, df))
check("t quadrature stability", abs(tcdf(critical, df, 8192) - tcdf(critical, df, 4096)) < 1e-12)

constant = sum(len(set(row["es"])) == 1 for row in groups.values())
transitions = Counter((row["es"][1], row["es"][2]) for row in groups.values())
changed = sum(a != b for a, b in transitions.elements())
check("66 unchanged and three final transitions", constant == 66 and changed == 3)
check("19 coder cells and five group endpoints differ", len(disagreements) == 19 and len(endpoint_disagreements) == 5)
coder2cells, coder2contrasts, _ = aggregate(coder2y)

# Compare only after constructing every audited quantity independently.
primary = json.loads((BASE / "results/analysis.json").read_text())["primary"]
for label, value in [("effect", float(effect)), ("se", se), ("df", df), ("two_sided_p", pvalue),
                     ("group_weighted_sensitivity", float(mean(group_contrasts)))]:
    check(f"matches primary {label}", abs(value - primary[label]) < 1e-11)
check("matches primary CI", max(abs(x - y) for x, y in zip(ci, primary["ci95"])) < 1e-11)
after = {name: sha256((DATA / name).read_bytes()).hexdigest() for name in source_names}
check("raw source hashes unchanged", before == after)

result = dict(
    methodology="Independent raw extraction; exact rational means/variances; Student-t numerical quadrature; no producer imports.",
    source_sha256=after, check_count=len(checks), checks=checks,
    primary=dict(effect_exact=str(effect), effect=float(effect), variance_exact=str(v), variance=float(v),
                 se=se, df=df, critical_t=critical, ci95=ci, two_sided_p=pvalue),
    cells={t: {key: [str(x) for x in value] if key == "values" else str(value) if isinstance(value, Q) else value
               for key, value in cell.items()} for t, cell in cells.items()},
    contrasts=list(map(str, contrasts)), group_weighted_contrasts=list(map(str, group_contrasts)),
    group_weighted_effect_exact=str(mean(group_contrasts)),
    constant_es_groups=constant, changed_second_to_third=changed,
    changed_groups={g: row["es"] for g, row in groups.items() if len(set(row["es"])) != 1},
    transitions={f"{a}->{b}": n for (a, b), n in sorted(transitions.items())},
    coder_disagreements=disagreements, endpoint_disagreements=endpoint_disagreements,
    coder2_descriptive_effect_exact=str(mean(coder2contrasts)),
    limits=["Allocation across sessions unverified", "No exact randomization guarantee", "Coder-2 result is post-result sensitivity",
            "Retained category does not establish active renewal", "Secondary token and penalty means not independently recomputed"],
)
out = BASE / "results/review_check.json"
out.write_text(json.dumps(result, indent=2) + "\n")
print(json.dumps({"checks_passed": len(checks), "effect": float(effect), "ci95": ci,
                  "coder_endpoint_disagreements": endpoint_disagreements, "output": str(out)}, indent=2))
