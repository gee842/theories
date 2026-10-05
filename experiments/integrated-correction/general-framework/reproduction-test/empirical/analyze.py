"""Frozen read-only reanalysis. Never modifies source workbooks."""
from collections import defaultdict, Counter
from pathlib import Path
import csv
import hashlib
import json
import math
import datetime
import subprocess

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
OUT = ROOT / "results"
OUT.mkdir(exist_ok=True)
TREATMENTS = {1: "CNN", 2: "FNN", 3: "NCN", 4: "NFN", 5: "PCN", 6: "PFN"}
ORIGINAL = {1: 4, 2: 2, 3: 1, 4: 3, 5: 5, 6: 6}
EXPECTED_PROTOCOL = "d99f55eb37b822aa5baec339bde0ea56e46c92260bde924d8b89ff2472fb4891"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_csv(path, rows):
    with path.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


assert digest(ROOT / "PROTOCOL.md") == EXPECTED_PROTOCOL
w = load_workbook(DATA / "Dataset2_Main.xlsx", read_only=True, data_only=True)
sheet = w["Data"]
iterator = sheet.iter_rows(values_only=True)
headers = list(next(iterator))
fields = ["SessID", "Group", "Player", "GraphTreatmentNum", "OrigTreatmentNum", "Facil"]
fields += [f"ES Rd{r}" for r in range(1, 4)]
fields += [f"TokensRd{r}" for r in range(1, 7)]
fields += [f"Trigg Penalty Rd{r}" for r in range(1, 7)]
indices = {f: headers.index(f) for f in fields}
people = []
for row_number, values in enumerate(iterator, start=2):
    if values[indices["Group"]] is not None:
        people.append({"source_row": row_number, **{f: values[i] for f, i in indices.items()}})
w.close()
by_group = defaultdict(list)
for row in people:
    by_group[int(row["Group"])].append(row)


def shared(rows, field):
    present = {r[field] for r in rows if r[field] is not None}
    assert len(present) <= 1, (field, present)
    return next(iter(present)) if present else None


groups = []
for group_id, rows in sorted(by_group.items()):
    assert len(rows) == 4
    assert {r["Player"] for r in rows} == {1, 2, 3, 4}
    treatment_num = int(shared(rows, "GraphTreatmentNum"))
    assert ORIGINAL[int(shared(rows, "OrigTreatmentNum"))] == treatment_num
    treatment = TREATMENTS[treatment_num]
    session = shared(rows, "SessID").split("_", 1)[1]
    es = [shared(rows, f"ES Rd{r}") for r in range(1, 4)]
    assert all(v is None or v in {0, 1, 2, 3} for v in es)
    y = None if any(v is None for v in es[1:]) else int(es[1] >= 2 and es[2] >= 2)
    rounds = range(1, 4) if treatment_num in {1, 2} else range(4, 7)
    token_values = [r[f"TokensRd{k}"] for r in rows for k in rounds]
    penalty_values = [r[f"Trigg Penalty Rd{k}"] for r in rows for k in rounds]
    assert all(isinstance(v, (int, float)) for v in token_values + penalty_values)
    groups.append({"group": group_id, "session": session, "treatment": treatment,
                   "facilitated": int("F" in treatment), "facil_source": shared(rows, "Facil"),
                   "es1": es[0], "es2": es[1], "es3": es[2], "retained_permission": y,
                   "constant_es": int(es[0] == es[1] == es[2]) if None not in es else None,
                   "mean_tokens_per_person_round": float(np.mean(token_values)),
                   "mean_penalties_per_group_round": float(sum(penalty_values) / 3),
                   "source_rows": ";".join(str(r["source_row"]) for r in rows)})
assert len(groups) == 72
missing = [r["group"] for r in groups if r["retained_permission"] is None]
if missing:
    (OUT / "missing.json").write_text(json.dumps({"missing_primary_groups": missing}, indent=2))
    raise RuntimeError("Primary ES missing; frozen endpoint cannot be silently replaced.")

by_session = defaultdict(list)
for row in groups:
    by_session[row["session"]].append(row)
sessions = []
for session, rows in sorted(by_session.items()):
    assert len({r["treatment"] for r in rows}) == 1
    sessions.append({"session": session, "treatment": rows[0]["treatment"],
                     "groups": len(rows), "mean_retained_permission": float(np.mean([r["retained_permission"] for r in rows]))})
assert len(sessions) == 31
assert all(r["facil_source"] == r["facilitated"] for r in groups)
cells = {}
for treatment in TREATMENTS.values():
    sr = [r["mean_retained_permission"] for r in sessions if r["treatment"] == treatment]
    gr = [r for r in groups if r["treatment"] == treatment]
    cells[treatment] = {"sessions": len(sr), "groups": len(gr), "session_mean": float(np.mean(sr)),
                        "session_variance": float(np.var(sr, ddof=1)),
                        "retained_groups": sum(r["retained_permission"] for r in gr),
                        "group_mean": float(np.mean([r["retained_permission"] for r in gr])),
                        "constant_es_groups": sum(r["constant_es"] for r in gr),
                        "mean_tokens_per_person_round": float(np.mean([r["mean_tokens_per_person_round"] for r in gr])),
                        "mean_penalties_per_group_round": float(np.mean([r["mean_penalties_per_group_round"] for r in gr]))}
pairs = [("CNN", "FNN"), ("NCN", "NFN"), ("PCN", "PFN")]
contrasts = [{"unfacilitated": c, "facilitated": f,
              "session_weighted_effect": cells[f]["session_mean"] - cells[c]["session_mean"],
              "group_weighted_effect": cells[f]["group_mean"] - cells[c]["group_mean"]} for c, f in pairs]
effect = float(np.mean([r["session_weighted_effect"] for r in contrasts]))
components = [cell["session_variance"] / cell["sessions"] / 9 for cell in cells.values()]
variance = sum(components)
se = math.sqrt(variance)
df = variance ** 2 / sum(v * v / (cell["sessions"] - 1) for v, cell in zip(components, cells.values()))
# The bundled Python has numpy/openpyxl but no SciPy. Use the existing research
# environment only for the missing Student-t CDF/quantile capability.
stats_python = "/private/tmp/meta-attractor-plot-env/bin/python"
stats = json.loads(subprocess.check_output([
    stats_python, "-c",
    "import sys,json; from scipy.stats import t; d,z=map(float,sys.argv[1:]); print(json.dumps([float(t.ppf(.975,d)),float(2*t.sf(abs(z),d))]))",
    str(df), str(effect / se)
], text=True))
critical, pvalue = stats
primary = {"effect": effect, "se": se, "df": df, "ci95": [effect - critical * se, effect + critical * se],
           "two_sided_p": pvalue,
           "group_weighted_sensitivity": float(np.mean([r["group_weighted_effect"] for r in contrasts]))}
transition_rows = []
for treatment in TREATMENTS.values():
    for before, after in [(1, 2), (2, 3)]:
        counts = Counter((r[f"es{before}"], r[f"es{after}"]) for r in groups if r["treatment"] == treatment)
        for a in range(4):
            for b in range(4):
                transition_rows.append({"treatment": treatment, "from_opportunity": before, "to_opportunity": after,
                                        "from_es": a, "to_es": b, "count": counts[(a, b)]})

# Audit authors' two independent coder columns without using them to choose an endpoint.
cw = load_workbook(DATA / "Dataset1_ChatCodes.xlsm", read_only=True, data_only=True)
cs = cw["Kappa (Interrater Reliability)"]
cr = cs.iter_rows(values_only=True)
ch = list(next(cr))
coder_counts = Counter()
for vals in cr:
    if vals[0] is None:
        continue
    gid = int(vals[0])
    main = next(r for r in groups if r["group"] == gid)
    for k in range(1, 4):
        a = vals[ch.index(f"Coder 1 ES Rd{k}")]
        b = vals[ch.index(f"Coder 2 ES Rd{k}")]
        coder_counts["pairs"] += 1
        coder_counts["disagreements"] += int(a != b)
        coder_counts["main_differs_coder1"] += int(main[f"es{k}"] != a)
        coder_counts["main_differs_coder2"] += int(main[f"es{k}"] != b)
cw.close()

provenance = {
    "analyzed_at_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(),
    "article": "https://doi.org/10.1371/journal.pone.0307832", "archive": "https://osf.io/v3drm/",
    "archive_license": "Not assigned in public OSF node metadata; article CC BY does not establish a separate data license",
    "protocol_sha256": digest(ROOT / "PROTOCOL.md"), "analysis_sha256": digest(Path(__file__)),
    "data_sha256": {name: digest(DATA / name) for name in ["Dataset1_ChatCodes.xlsm", "Dataset2_Main.xlsx", "CODEBOOK.pdf", "Appendix.docx", "node.json"]},
    "source_ranges": {f: f"Data!{get_column_letter(i+1)}2:{get_column_letter(i+1)}289" for f, i in indices.items()},
    "coder_range": "Kappa (Interrater Reliability)!A1:M73"
}
result = {"primary": primary, "cells": cells, "contrasts": contrasts,
          "constant_es_groups": sum(r["constant_es"] for r in groups), "total_groups": len(groups),
          "changed_second_to_third": sum(r["es2"] != r["es3"] for r in groups),
          "coder_audit": dict(coder_counts), "provenance": provenance}
write_csv(OUT / "groups.csv", groups)
write_csv(OUT / "sessions.csv", sessions)
write_csv(OUT / "transitions.csv", transition_rows)
(OUT / "analysis.json").write_text(json.dumps(result, indent=2))
print(json.dumps({k: v for k, v in result.items() if k != "provenance"}, indent=2))
