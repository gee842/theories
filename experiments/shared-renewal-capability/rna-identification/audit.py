#!/usr/bin/env python3
"""Read-only secondary audit of published RNA inputs; no model fitting or simulation."""
from collections import defaultdict
from fractions import Fraction as Q
from pathlib import Path
from statistics import mean
import hashlib
import json
import math
import re

ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
SOURCE = ROOT / 'sources/24-rna-network-evidence-2026-10-05'
ADD = SOURCE / 'identification-addition'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    from openpyxl import load_workbook
    manifests = [(SOURCE, json.loads((SOURCE / 'MANIFEST.json').read_text())),
                 (ADD, json.loads((ADD / 'MANIFEST.json').read_text()))]
    before = {}
    for directory, manifest in manifests:
        for item in manifest['sources']:
            path = directory / item['file']
            assert digest(path) == item['sha256'], path
            before[str(path.relative_to(ROOT))] = digest(path)

    path = ADD / 'Mizuuchi-2022-source-data.xlsx'
    wb = load_workbook(path, read_only=True, data_only=False)
    inventory = [{'sheet': s.title, 'rows': s.max_row, 'columns': s.max_column}
                 for s in wb]
    assert [x['sheet'] for x in inventory] == [
        'Fig. 4', 'Fig. S1', 'Fig. S5', 'Fig. S13', 'Fig. S14', 'Fig. S19']
    sheet = wb['Fig. 4']
    grouped = defaultdict(list)
    for row in range(10, 58):
        donor, target, own, partner = [sheet.cell(row, c).value for c in (14, 15, 16, 17)]
        if donor:
            assert isinstance(own, (int, float))
            assert partner is None or isinstance(partner, (int, float))
            grouped[donor, target].append({'row': row, 'donor_fold': own,
                                           'target_fold': partner})
    names = ['HL1-228', 'HL2-228', 'HL3-228', 'PL2-228', 'PL3-228']
    # Numeric table, visually transcribed from original SI p29, Fig S23.
    # Rows are replicase donors j, columns RNA targets i, matching the notebook.
    printed = [[2.615577, 0, 0, 0, 3.372959],
               [1.419954, 2.79044, 1.422017, 5.711838, 3.040918],
               [0, 0, 0.438937, 0, 3.058117],
               [0, 0, 0, 0, 0], [0, 0, 0, 0, 0]]
    matrix, reconstruction = [], []
    for donor in names[:3]:
        self_rows = grouped[donor, None]
        alpha = mean(x['donor_fold'] for x in self_rows)
        values = []
        for target in names:
            if donor == target:
                rows, beta, gamma, reff = self_rows, None, None, alpha
                active = alpha > 1.5
            else:
                rows = grouped[donor, target]
                beta = mean(x['donor_fold'] for x in rows)
                gamma = mean(x['target_fold'] for x in rows)
                reff = alpha * gamma / beta
                active = gamma > 1.5
            k = math.log(reff) if active else 0.0
            values.append(k)
            reconstruction.append(dict(
                donor=donor, target=target, replicate_rows=rows,
                self_rows=[x['row'] for x in self_rows], alpha=alpha, beta=beta,
                gamma=gamma, reff=reff, target_mean_above_1_5=active,
                log_rate=k, cell_columns={'donor': 'N', 'target': 'O',
                                         'donor_fold': 'P', 'target_fold': 'Q'}))
        matrix.append(values)
    matrix += [[0.0]*5, [0.0]*5]  # Parasites do not encode replicase; model zeros.
    errors = [abs(matrix[j][i] - printed[j][i]) for j in range(5) for i in range(5)]
    assert max(errors) < 5e-7
    assert grouped['HL1-228', None][0]['row'] == 10
    assert sheet['J23'].value == 'sample loss'  # Earlier-round missing value, unused.
    assert [wb['Fig. S19'][f'B{r}'].value for r in (11, 12, 13)] == ['HL1-158']*3

    kh, kp = matrix[1][1], matrix[1][3]
    ratio = kp / kh
    D, C, h0, p0 = 5, 300, 10, 10
    required = D*h0 + D**ratio*p0
    alpha_values = [x['donor_fold'] for x in grouped['HL2-228', None]]
    pair_rows = grouped['HL2-228', 'PL2-228']
    beta_values = [x['donor_fold'] for x in pair_rows]
    gamma_values = [x['target_fold'] for x in pair_rows]
    r_min = 1 + math.log(min(gamma_values)/max(beta_values))/math.log(max(alpha_values))
    r_max = 1 + math.log(max(gamma_values)/min(beta_values))/math.log(min(alpha_values))
    assert required > C and r_min < 2 < r_max

    # Exact Michaelis-kinetics nonidentifiability witness from unambiguous
    # original-host kcat/Km component values, not invented fitted on/off rates.
    kcat, km = Q(22, 100), Q(76, 100)
    enzyme_witness = []
    for kon in (Q(1), Q(2)):
        koff = kon*km-kcat
        assert koff > 0 and (koff+kcat)/kon == km
        enzyme_witness.append(dict(kon=str(kon), koff=str(koff),
                                   kcat=str(kcat), Km=str(km),
                                   bound_lifetime=str(1/(koff+kcat))))

    notebook = json.loads((ADD / 'Serial transfer of the five RNAs.ipynb').read_text())
    code = '\n'.join(''.join(c['source']) for c in notebook['cells'] if c['cell_type'] == 'code')
    external_csv = re.findall(r"loadtxt\('([^']+)'", code)
    assert external_csv == ['211017.csv']
    assert 'reaction_time = 2' in code and 'Cn = 300' in code and 'D=5' in code
    assert 'Fnumber=int(M*1.3)' in code
    assert before == {rel: digest(ROOT / rel) for rel in before}
    result = dict(
        status='passed; descriptive reconstruction and conditional algebra only',
        source_sha256_before_and_after=before, checker_sha256=digest(Path(__file__)),
        workbook_inventory=inventory, round228_assay_rows=48,
        matrix_axis_convention='rows donor j, columns target i',
        matrix_names=names, reconstructed_matrix=matrix, printed_S23_matrix=printed,
        maximum_absolute_rounding_difference=max(errors), reconstruction=reconstruction,
        missing_notebook_input=external_csv,
        missing_input_scope='Absent from pinned author repository tree (README and notebook only); no substitute CSV represented as original.',
        notebook_settings=dict(compartments=300000, capacity=300, dilution=5,
                               fusion_division_per_compartment=1.3, reaction_time_code_units=2),
        source_ambiguities=['Fig S19 B11:B13 label HL1-158 while SI caption says round228.',
                            'Fig4 J23 sample loss retained as text; not part of round228 reconstruction.',
                            'Bansho2016 evolved-host efficiency/ratio printed inconsistency remains unresolved.'],
        joint_renewal_boundary=dict(
            model='Two-RNA restriction of published effective ODE, common saturation, no mutation/degradation',
            host='HL2-228', parasite='PL2-228', kh=kh, kp=kp, r=ratio,
            dilution=D, capacity=C, initial_host=h0, initial_parasite=p0,
            required_capacity_for_joint_fivefold=required,
            critical_equal_initial_occupancy=C/(D+D**ratio),
            condition='D*h0 + D**r*p0 < C; strict for finite time when r>1',
            observed_extrema_sensitivity_r=[r_min, r_max],
            sensitivity_scope='Range from component observed extrema; not a confidence interval or paired resampling'),
        exact_michaelis_nonidentifiability_witness=enzyme_witness,
        scope='No original notebook execution, fitting, stochastic data reproduction, held-out validation, or CF mapping.')
    (HERE / 'audit.json').write_text(json.dumps(result, indent=2)+'\n')
    print(json.dumps({k: result[k] for k in ['status', 'maximum_absolute_rounding_difference',
                                          'joint_renewal_boundary']}, indent=2))


if __name__ == '__main__':
    main()
