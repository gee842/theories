#!/usr/bin/env python3
"""Reproduce descriptive checks without modifying the preserved Ceroni workbook."""

from __future__ import annotations

import argparse
import ast
import hashlib
import json
import math
from pathlib import Path
import platform
import statistics

import openpyxl


REPO = Path(__file__).resolve().parents[3]
EVIDENCE = REPO / (
    "_process/research/2026-10-05-support-and-shared-generativity/"
    "completion/evidence/ceroni-2018"
)
SOURCE = EVIDENCE / "41592_2018_BFnmeth4635_MOESM5_ESM.xlsx"
EXPECTED_SHA256 = "ac386fc130aec5f79bc3407961710df421a4ac9377faabf8f43bf2a5298c47ed"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def triplicate(sheet, coordinate: str) -> dict:
    value = sheet[coordinate].value
    values = ast.literal_eval(value) if isinstance(value, str) else value
    if not isinstance(values, list) or len(values) != 3:
        raise ValueError(f"Expected three observations at {sheet.title}!{coordinate}")
    if not all(isinstance(x, (float, int)) and math.isfinite(x) for x in values):
        raise ValueError(f"Non-numeric observation at {sheet.title}!{coordinate}")
    return {
        "sheet": sheet.title,
        "cell": coordinate,
        "source_cell_value": value,
        "values": values,
        "n": len(values),
        "mean": statistics.mean(values),
        "sample_standard_deviation": statistics.stdev(values),
    }


def compare_mean(raw_sheet, raw_cell, summary_sheet, summary_cell) -> dict:
    result = triplicate(raw_sheet, raw_cell)
    summary = summary_sheet[summary_cell].value
    difference = result["mean"] - summary
    result.update({
        "summary_sheet": summary_sheet.title,
        "summary_cell": summary_cell,
        "supplied_summary_value": summary,
        "recomputed_minus_supplied": difference,
        "agrees_within_absolute_1e_8": abs(difference) <= 1e-8,
    })
    return result


def compare_sheets(left, right) -> dict:
    """Exact cached-value comparison, including strings and empty cells, not styles."""
    rows = max(left.max_row, right.max_row)
    columns = max(left.max_column, right.max_column)
    left_values = list(left.iter_rows(max_row=rows, max_col=columns, values_only=True))
    right_values = list(right.iter_rows(max_row=rows, max_col=columns, values_only=True))
    different = []
    difference_count = 0
    for row_idx, (left_row, right_row) in enumerate(zip(left_values, right_values), 1):
        for col_idx, (left_value, right_value) in enumerate(zip(left_row, right_row), 1):
            if left_value != right_value:
                difference_count += 1
                if len(different) < 4:
                    different.append({
                        "cell": f"{openpyxl.utils.get_column_letter(col_idx)}{row_idx}",
                        "left": left_value,
                        "right": right_value,
                    })
    return {
        "left_sheet": left.title,
        "right_sheet": right.title,
        "left_dimensions": [left.max_row, left.max_column],
        "right_dimensions": [right.max_row, right.max_column],
        "comparison_dimensions": [rows, columns],
        "exact_cell_value_equality": difference_count == 0,
        "different_cell_count": difference_count,
        "first_differences": different,
    }


def audit() -> dict:
    before = digest(SOURCE)
    if before != EXPECTED_SHA256:
        raise ValueError("Preserved source hash differs from the audited source; refusing to proceed")
    manifest = json.loads((EVIDENCE / "MANIFEST.json").read_text())
    manifest_entry = next(x for x in manifest["files"] if x["path"] == SOURCE.name)
    if manifest_entry["sha256"] != EXPECTED_SHA256:
        raise ValueError("Source manifest does not match the fixed audit hash")

    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    try:
        batch = workbook["Fig S14a raw data"]
        batch_rows = []
        for column in "CDEF":
            values = {str(row): batch[f"{column}{row}"].value for row in [1, 6, 13, 20]}
            ratios = {
                "od700": values["6"] / batch["C6"].value,
                "red_fluorescence_per_od": values["13"] / batch["C13"].value,
                "total_red_fluorescence": values["20"] / batch["C20"].value,
            }
            residual = ratios["od700"] * ratios["red_fluorescence_per_od"] - ratios["total_red_fluorescence"]
            batch_rows.append({
                "column": column,
                "source_cells": {f"{column}{row}": value for row, value in values.items()},
                "ratios_to_column_c": ratios,
                "ratio_product_minus_total_ratio": residual,
                "ratio_product_agrees_within_absolute_1e_12": abs(residual) <= 1e-12,
            })

        raw_b = workbook["Fig S8 Raw Data B"]
        mean_b = workbook["Fig S8 Mean and Std values B"]
        induced = []
        for column, summary_column in [("E", "G"), ("I", "O"), ("J", "Q")]:
            induced.append({
                "raw_column": column,
                "source_heading": raw_b[f"{column}1"].value,
                "growth_rate_per_hour": compare_mean(raw_b, f"{column}5", mean_b, f"{summary_column}6"),
                "gfp_rate_per_od_per_hour": compare_mean(raw_b, f"{column}33", mean_b, f"{summary_column}35"),
            })
        h2_growth = induced[-1]["growth_rate_per_hour"]["mean"]
        for item in induced:
            item["growth_difference_from_h2_per_hour"] = item["growth_rate_per_hour"]["mean"] - h2_growth

        raw_a = workbook["Fig S8 Raw Data A"]
        mean_a = workbook["Fig S8 Mean and Std values A"]
        discrepancies = [
            compare_mean(raw_a, "D32", mean_a, "E34"),
            compare_mean(raw_a, "G5", mean_a, "K6"),
        ]
        mean_comparison = compare_sheets(workbook["Fig S9 mean & stdev"], workbook["Fig S13 mean & stdev"])
        raw_comparison = compare_sheets(workbook["Fig S9 raw data"], workbook["Fig S13 raw data"])

        checks = {
            "s14a_headers_are_as_reported": [x["source_cells"][f"{x['column']}1"] for x in batch_rows] == ["H2", "FB5", "FB", "FB2"],
            "s14a_non_h2_columns_have_more_od_less_per_od_and_more_total": all(
                x["ratios_to_column_c"]["od700"] > 1
                and x["ratios_to_column_c"]["red_fluorescence_per_od"] < 1
                and x["ratios_to_column_c"]["total_red_fluorescence"] > 1
                for x in batch_rows[1:]
            ),
            "s14a_product_identities_hold": all(x["ratio_product_agrees_within_absolute_1e_12"] for x in batch_rows),
            "selected_s8_induced_means_match_summaries": all(
                x[k]["agrees_within_absolute_1e_8"]
                for x in induced
                for k in ["growth_rate_per_hour", "gfp_rate_per_od_per_hour"]
            ),
            "both_selected_s8_uninduced_means_disagree": all(not x["agrees_within_absolute_1e_8"] for x in discrepancies),
            "s9_s13_mean_sheets_are_both_81_rows": mean_comparison["left_dimensions"][0] == mean_comparison["right_dimensions"][0] == 81,
            "s9_s13_mean_sheets_have_exactly_equal_values": mean_comparison["exact_cell_value_equality"],
            "s9_s13_raw_sheets_have_unequal_values": not raw_comparison["exact_cell_value_equality"],
        }
        result = {
            "audit_type": "Exploratory secondary descriptive source-data audit; not validation or replication",
            "runtime": {"python": platform.python_version(), "openpyxl": openpyxl.__version__},
            "source": {
                "repository_relative_path": SOURCE.relative_to(REPO).as_posix(),
                "url": manifest_entry["url"],
                "expected_sha256": EXPECTED_SHA256,
                "sha256_before": before,
                "bytes": SOURCE.stat().st_size,
                "open_mode": "openpyxl read_only=True, data_only=True; workbook never saved",
            },
            "s14a_24_hour": {
                "sheet": batch.title,
                "time_cell": "B6",
                "time_hours": batch["B6"].value,
                "columns": batch_rows,
                "non_h2_total_ratio_min_max": [
                    min(x["ratios_to_column_c"]["total_red_fluorescence"] for x in batch_rows[1:]),
                    max(x["ratios_to_column_c"]["total_red_fluorescence"] for x in batch_rows[1:]),
                ],
                "interpretation": "Column labels are retained without repair. The paper figure assigns the non-H2 trajectories different feedback labels. Ratios are descriptive, not confidence bounds; per-OD values share the yield and OD data.",
            },
            "s8_induced_one_hour": induced,
            "s8_uninduced_raw_vs_summary_discrepancies": discrepancies,
            "s9_vs_s13_mean_sheets": mean_comparison,
            "s9_vs_s13_raw_sheets": raw_comparison,
            "checks": checks,
        }
    finally:
        workbook.close()
        after = digest(SOURCE)
        if after != before:
            raise RuntimeError("Source workbook hash changed during audit")

    result["source"]["sha256_after"] = after
    result["checks"]["source_unchanged_and_matches_fixed_hash"] = before == after == EXPECTED_SHA256
    result["all_expected_checks_passed"] = all(result["checks"].values())
    if not result["all_expected_checks_passed"]:
        raise AssertionError(json.dumps(result["checks"], indent=2))
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path(__file__).with_name("ceroni_audit.json"))
    args = parser.parse_args()
    if args.output.resolve() == SOURCE.resolve():
        raise ValueError("Output must not replace the source workbook")
    result = audit()
    args.output.write_text(json.dumps(result, indent=2, allow_nan=False) + "\n")
    print(f"All expected audit checks passed; source unchanged. Results: {args.output}")


if __name__ == "__main__":
    main()
