"""Read-only source semantics audit; never imports the primary analysis.

Requires openpyxl for unchanged cached/formula workbook access. Writes only its
own receipt. It classifies the 39 existing source_summary_audit discrepancies;
it does not change a source, estimand, result, participant set, or model.
"""
from pathlib import Path
from fractions import Fraction
from collections import Counter
import csv
import hashlib
import json
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SOURCE = ROOT / 'sources/30-recipient-harm-valuation-2026-10-06'
RAW = SOURCE / 'original'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    trial_rows = list(csv.reader((RAW / 'TrialByTrialChoicesMainPaper.csv').open(encoding='utf-8-sig', newline='')))
    trial = {int(r[1]): r for r in trial_rows[2:] if r[1]}
    summary = {int(r['Animal Number']): r for r in csv.DictReader((RAW / 'dataMainPaper.csv').open(encoding='utf-8-sig', newline='')) if r['Animal Number']}
    formula = openpyxl.load_workbook(RAW / 'dataMainPaper.xlsx', read_only=True, data_only=False)
    cached = openpyxl.load_workbook(RAW / 'dataMainPaper.xlsx', read_only=True, data_only=True)
    primary = json.loads((HERE / 'results.json').read_text())
    discrepancies = primary['source_summary_audit']['discrepancies']
    expected = {(i, 'nonperformance_total') for i in range(47, 79)} | {(37, 'AverageShock')} | {(i, field) for i in (70, 73, 78) for field in ('Shock 3', 'AverageShock')}
    assert len(discrepancies) == len(expected) == 39
    assert {(x['actor'], x['field']) for x in discrepancies} == expected
    assert all(trial[i][2] == summary[i]['Condition'] for i in trial)
    findings = []

    for item in discrepancies:
        i, field = item['actor'], item['field']
        record = {'actor': i, 'field': field, 'original_discrepancy': item,
                  'summary_group': summary[i]['Group'], 'sbl': summary[i]['SBL#']}
        if field == 'nonperformance_total':
            cell = f'CJ{i+2}'
            assert trial[i][87] == ''
            assert formula['Choice'][cell].value is None
            assert cached['Choice'][cell].value is None
            zero_counts = [trial[i][start:start+20].count('0') for start in (3, 23, 43, 63)]
            assert [int(v) for v in trial[i][83:87]] == zero_counts
            assert item['computed'] == sum(zero_counts)
            record.update(category='omitted_total_summary', csv_total='', workbook_cell='Choice!'+cell,
                          workbook_formula_value=None, workbook_cached_value=None,
                          phase_zero_counts=zero_counts, reconstructed_total=sum(zero_counts),
                          interpretation='Absent summary cell, not a contradictory reported count.')
        elif i == 37:
            assert trial[i][63:83] == ['0'] * 20
            assert summary[i]['Shock 3'] == 'NaN'
            assert formula['data']['J38'].value == 'NaN'
            assert formula['data']['K38'].value == '=AVERAGE(H38:J38)'
            available = [cached['data'][c].value for c in ('H38', 'I38')]
            assert available == [95, 100]
            assert cached['data']['K38'].value == sum(available)/2 == 97.5
            assert item['computed'] is None
            record.update(category='spreadsheet_average_ignores_undefined_session',
                          shock_summary_cells={'data!H38': 95, 'data!I38': 100, 'data!J38': 'NaN'},
                          average_cell='data!K38', average_formula='=AVERAGE(H38:J38)', average_cached=97.5,
                          shock3_trial_counts={'0': 20},
                          interpretation='Workbook mean over two numeric sessions differs from a mean requiring all three conditional choices to be defined.')
        else:
            phases = []
            for name, start in zip(('Shock 1', 'Shock 2', 'Shock 3'), (23, 43, 63)):
                counts = Counter(trial[i][start:start+20])
                assert len(trial[i][start:start+20]) == 20
                scheduled = Fraction(100*counts['2'], 20)
                completed = Fraction(100*counts['2'], counts['1']+counts['2'])
                assert Fraction(summary[i][name]) == scheduled
                phases.append({'phase': name, 'counts': dict(counts), 'scheduled_percentage': str(scheduled), 'completed_response_percentage': str(completed)})
            assert cached['data'].cell(i+1, 10).value == int(phases[2]['scheduled_percentage'])
            assert formula['data'].cell(i+1, 11).value == f'=AVERAGE(H{i+1}:J{i+1})'
            avg_scheduled = sum(Fraction(p['scheduled_percentage']) for p in phases) / 3
            avg_completed = sum(Fraction(p['completed_response_percentage']) for p in phases) / 3
            assert abs(Fraction(summary[i]['AverageShock'])-avg_scheduled) <= Fraction(1, 200)
            target = Fraction(phases[2]['completed_response_percentage']) if field == 'Shock 3' else avg_completed
            assert Fraction(item['computed']) == target
            record.update(category='scheduled_versus_completed_response_denominator',
                          phases=phases, average_scheduled_percentage=str(avg_scheduled),
                          average_completed_response_percentage=str(avg_completed),
                          workbook_average_cell=f'data!K{i+1}',
                          workbook_average_formula=f'=AVERAGE(H{i+1}:J{i+1})',
                          interpretation='Exact arithmetic match to code2 divided by 20 in the supplied summaries; this identifies a denominator difference, not the original analyst intention. A blank does not become a known code.')
        findings.append(record)

    # A second denominator distinction in the provided total column: its original
    # formulas cover shock phases only. All observed baseline counts are zero,
    # so this has no numerical effect in these released rows.
    total_formulas = {}
    for i in range(1, 47):
        cell = f'CJ{i+2}'
        expected_formula = f'=SUM(CG{i+2}:CI{i+2})'
        assert formula['Choice'][cell].value == expected_formula
        total_formulas[str(i)] = expected_formula
    assert all(trial[i][3:23].count('0') == 0 for i in trial)
    assert not any(x['actor'] in range(1, 27) for x in discrepancies)
    tolerance = Fraction(11, 1000)
    rounding_only_fields = []
    for x in primary['source_summary_audit']['comparisons']:
        if x.get('absolute_difference') is not None and x['agree']:
            difference = Fraction(x['absolute_difference'])
            assert difference <= tolerance
            if difference:
                rounding_only_fields.append({'actor': x['actor'], 'field': x['field'], 'absolute_difference': str(difference)})

    source_inputs = [RAW/'TrialByTrialChoicesMainPaper.csv', RAW/'dataMainPaper.csv', RAW/'dataMainPaper.xlsx', RAW/'hernandez2020.pdf', SOURCE/'manifest.json', HERE/'results.json', HERE/'PROTOCOL.md', Path(__file__)]
    receipt = {'status': 'pass', 'scope': 'Exhaustive semantics classification of the 39 existing discrepancy flags. No redefinition of primary estimates.',
               'discrepancies_classified': len(findings), 'categories': dict(Counter(x['category'] for x in findings)),
               'findings': findings, 'first46_total_formulas_cover_shock_phases_only': total_formulas,
               'all78_baseline_zero_counts_are_zero': True,
               'primary26_male_actor_discrepancy_count': 0,
               'ordinary_rounding_fields_already_within_primary_tolerance': rounding_only_fields,
               'methods_support': {'PDF18_e3': '20 free choices after four forced samples; failed free trials advance.', 'PDF20_e5': 'Forced trials excluded from all analyses; failed trials separately quantified.', 'PDF21_e6': 'Sshock defined as average proportion across three shock sessions; no explicit all-nonresponse convention.'},
               'input_sha256': {str(p.relative_to(ROOT)): sha(p) for p in source_inputs}}
    destination = HERE / 'source-summary-audit.json'
    destination.write_text(json.dumps(receipt, indent=2)+'\n')
    print(json.dumps({'status': 'pass', 'discrepancies': 39, 'categories': receipt['categories'], 'receipt_sha256': sha(destination)}))


if __name__ == '__main__':
    main()
