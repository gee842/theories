#!/usr/bin/env python3
"""Read-only source extraction and frozen paired binned-data analysis.

Requires openpyxl and mpmath. Does not import the independent checker or mutate
any workbook. Student-t quantiles use a regularized-beta CDF and bisection.
"""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
from functools import lru_cache
import hashlib
import json
import math
from pathlib import Path
import platform
import statistics

import mpmath as mp
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SOURCE = ROOT / 'sources/29-motor-learning-cost-2026-10-06'
WORKBOOK = SOURCE / 'cost_of_error_data.xlsx'
EXPECTED_SOURCE = '0cd6ae5b0f159531f2442f49a474c5ecfda23a1fcf0965598da99df5d84891cc'
EXPECTED_PROTOCOL = '070940c7933c4aaf4115254b931230eeb50e7c08bffd214a20d3fe4ad3c28594'
BLOCKS = [
    ('EXP1', 'ASYMPTOTE', 'LARGE COST', 3),
    ('EXP1', 'ASYMPTOTE', 'SMALL COST', 27),
    ('EXP2', 'SWITCH', 'LARGE COST', 51),
    ('EXP2', 'SWITCH', 'SMALL COST', 75),
    ('CONT1', None, 'LOW COH', 147),
    ('CONT1', None, 'HIGH COH', 171),
]
WINDOWS = {
    'pre': (801, 931, 'AZ', 'BG', 8),
    'early': (971, 1100, 'BJ', 'BQ', 8),
    'late': (1101, 1250, 'BR', 'BZ', 9),
}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(name, value):
    (HERE / name).write_text(json.dumps(value, indent=2, allow_nan=False) + '\n')


def numeric(value, address):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f'Expected finite numeric cell at {address}: {value!r}')
    return float(value)


@lru_cache(maxsize=None)
def t975(df):
    """Positive Student-t .975 quantile, 50-digit beta CDF, 180 bisections."""
    with mp.workdps(50):
        nu = mp.mpf(df)
        def cdf(t):
            x = nu / (nu + t*t)
            return 1 - mp.betainc(nu/2, mp.mpf('.5'), 0, x, regularized=True)/2
        low, high = mp.mpf(0), mp.mpf(1)
        while cdf(high) < mp.mpf('.975'):
            high *= 2
        for _ in range(180):
            mid = (low + high)/2
            if cdf(mid) < mp.mpf('.975'):
                low = mid
            else:
                high = mid
        return float((low+high)/2)


def summary(values):
    n = len(values)
    assert n > 1
    mean = statistics.fmean(values)
    sd = statistics.stdev(values)
    se = sd/math.sqrt(n)
    critical = t975(n-1)
    return dict(n=n, mean=mean, sample_sd=sd, standard_error=se,
                df=n-1, t_critical=critical,
                ci95=[mean-critical*se, mean+critical*se])


def main(inventory_only=False):
    assert sha(WORKBOOK) == EXPECTED_SOURCE
    assert sha(HERE / 'PROTOCOL.md') == EXPECTED_PROTOCOL
    manifest = json.loads((SOURCE / 'MANIFEST.json').read_text())
    for filename, record in manifest['files'].items():
        assert sha(SOURCE / filename) == record['sha256'], filename
        assert (SOURCE / filename).stat().st_size == record['bytes'], filename
    before = {p.name: sha(p) for p in SOURCE.iterdir() if p.is_file()}
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=False)
    assert [(s.title, s.max_row, s.max_column) for s in wb] == [
        ('ADAPTATION',338,99), ('DECISION',146,78)]
    sheet = wb['ADAPTATION']
    inventory = dict(schema=1, source_sha256=EXPECTED_SOURCE,
                     protocol_sha256=EXPECTED_PROTOCOL,
                     sheets=[dict(name=s.title,rows=s.max_row,columns=s.max_column) for s in wb],
                     blocks=[], missing_values=0, participant_exclusions=[])
    extracted = []
    means = {}
    for experiment, subtitle, label, header in BLOCKS:
        expected = [experiment, subtitle, label, 'ADAPTATION']
        assert [sheet.cell(header,c).value for c in range(1,5)] == expected
        trialrow, firstrow, lastrow = header+2, header+3, header+22
        ids = [sheet.cell(r,2).value for r in range(firstrow,lastrow+1)]
        assert len(ids) == len(set(ids)) == 20 and ids == list(range(1,21))
        codes = [sheet.cell(r,1).value for r in range(firstrow,lastrow+1)]
        assert set(codes) == {1 if label in ('LARGE COST','LOW COH') else 2}
        block = dict(experiment=experiment, source_label=label, subtitle=subtitle,
                     header_row=header, block_row=header+1, trial_row=trialrow,
                     participant_rows=list(range(firstrow,lastrow+1)), participant_ids=ids,
                     condition_codes=codes, windows={}, guarded_out=[])
        for col in ('BH','BI'):
            c=column_index_from_string(col)
            coordinate=numeric(sheet.cell(trialrow,c).value, f'{col}{trialrow}')
            assert abs(coordinate-951) < 20
            block['guarded_out'].append(dict(column=col,trial_cell=f'{col}{trialrow}',mean_trial=coordinate))
        for window,(start,end,first,last,n) in WINDOWS.items():
            cols=list(range(column_index_from_string(first),column_index_from_string(last)+1))
            assert len(cols) == n
            # Select from all published coordinates independently of the expected columns.
            selected=[c for c in range(4,sheet.max_column+1)
                      if isinstance(sheet.cell(trialrow,c).value,(int,float))
                      and not isinstance(sheet.cell(trialrow,c).value,bool)
                      and start <= sheet.cell(trialrow,c).value <= end]
            assert selected == cols, (experiment,label,window,selected)
            coordinates=[numeric(sheet.cell(trialrow,c).value,sheet.cell(trialrow,c).coordinate) for c in cols]
            block['windows'][window]=dict(columns=[get_column_letter(c) for c in cols],
                mean_trial_coordinates=coordinates, trial_cells=[sheet.cell(trialrow,c).coordinate for c in cols],
                block_labels=[sheet.cell(header+1,c).value for c in cols], count=n)
            if experiment == 'CONT1':
                continue
            for row,pid in zip(range(firstrow,lastrow+1),ids):
                cells=[sheet.cell(row,c) for c in cols]
                values=[numeric(cell.value,cell.coordinate) for cell in cells]
                extracted.append(dict(experiment=experiment,source_label=label,
                    participant_id=pid,window=window,source_row=row,sheet=sheet.title,
                    cells=[cell.coordinate for cell in cells],values=values,
                    mean_trial_coordinates=coordinates))
                if not inventory_only:
                    means[(experiment,pid,label,window)]=statistics.fmean(values)
        inventory['blocks'].append(block)
    wb.close()
    save('inventory.json', inventory)
    save('extracted.json', dict(schema=1,source_sha256=EXPECTED_SOURCE,
                              protocol_sha256=EXPECTED_PROTOCOL,records=extracted))
    if inventory_only:
        print('Inventory and extraction saved; no contrasts calculated.')
        return
    participants=[]
    for exp in ('EXP1','EXP2'):
        for pid in range(1,21):
            wm={label:{w:means[(exp,pid,label,w)] for w in WINDOWS}
                for label in ('LARGE COST','SMALL COST')}
            changes={label:{w:wm[label][w]-wm[label]['pre'] for w in ('early','late')}
                     for label in wm}
            participants.append(dict(experiment=exp,participant_id=pid,window_means=wm,
                changes=changes,contrast={w:changes['SMALL COST'][w]-changes['LARGE COST'][w]
                                          for w in ('early','late')}))
    groups={}
    for exp in ('EXP1','EXP2'):
        pp=[p for p in participants if p['experiment']==exp]
        groups[exp]={w:dict(
            source_small_change=summary([p['changes']['SMALL COST'][w] for p in pp]),
            source_large_change=summary([p['changes']['LARGE COST'][w] for p in pp]),
            paired_contrast=summary([p['contrast'][w] for p in pp])) for w in ('early','late')}
    between={}
    for w in ('early','late'):
        one,two=[groups[exp][w]['paired_contrast'] for exp in ('EXP1','EXP2')]
        v1,v2=one['sample_sd']**2/one['n'],two['sample_sd']**2/two['n']
        se=math.sqrt(v1+v2)
        df=(v1+v2)**2/(v1*v1/(one['n']-1)+v2*v2/(two['n']-1))
        crit=t975(df)
        difference=two['mean']-one['mean']
        between[w]=dict(mean_difference_EXP2_minus_EXP1=difference,standard_error=se,
            df=df,t_critical=crit,ci95=[difference-crit*se,difference+crit*se],n_EXP1=one['n'],n_EXP2=two['n'])
    save('results.json',dict(schema=1,source_sha256=EXPECTED_SOURCE,protocol_sha256=EXPECTED_PROTOCOL,
        units='source baseline-relative amplitude change, degrees',
        contrast_definition='(SMALL COST post - pre) - (LARGE COST post - pre)',
        interpretation='literal source-series contrast; opposed-switch reading conditional on fixed-side labels',
        participants=participants,groups=groups,between_cohorts=between,
        validation=dict(cohorts=2,paired_participants=40,window_records=len(extracted),
                        analyzed_numeric_cells=sum(len(r['values']) for r in extracted),
                        exclusions=[],imputations=0)))
    after={p.name:sha(p) for p in SOURCE.iterdir() if p.is_file()}
    assert before == after, 'source directory changed during analysis'
    save('verification.json',dict(status='producer_passed_pending_independent_comparison',
        timestamp_utc=datetime.now(timezone.utc).isoformat(),
        source_unchanged=True,source_sha256=EXPECTED_SOURCE,protocol_sha256=EXPECTED_PROTOCOL,
        program_sha256=sha(Path(__file__)),
        versions=dict(python=platform.python_version(),openpyxl=openpyxl.__version__,mpmath=mp.__version__),
        statistical_method='Student-t CDF from regularized incomplete beta, 50-digit arithmetic, 180-step bisection',
        outputs={name:sha(HERE/name) for name in ('inventory.json','extracted.json','results.json')},
        checks=dict(source_files=len(manifest['files']),blocks=6,analyzed_blocks=4,
                    paired_participants=40,finite_analyzed_values=2000,missing=0,exclusions=0)))
    print('Producer completed: 40 paired participants, 2000 source values; outputs saved without workbook mutation.')


if __name__ == '__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--inventory-only',action='store_true')
    main(parser.parse_args().inventory_only)
